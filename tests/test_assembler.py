"""Tests del assembler, docx builder y formateador APA (sin mocks — CLAUDE.md)."""

import socket

import pytest
from docx import Document

from ic.domain.assembler import PilotAssembler
from ic.domain.chains.chain_01_descripcion import build_chain_01
from ic.domain.chains.chain_02_gics import build_chain_02
from ic.domain.chains.chain_03_benchmarking import build_chain_03
from ic.domain.use_cases.classify_gics import ClassifyGICS
from ic.domain.use_cases.generate_benchmarking import GenerateBenchmarking
from ic.domain.use_cases.generate_empresa_description import GenerateEmpresaDescription
from ic.infrastructure.citation.apa_formatter import (
    build_references_section,
    extract_citation_metadata,
    format_inline,
    format_reference,
)
from ic.infrastructure.docx_builder import DocxBuilder
from ic.infrastructure.llm.gateway import build_llm
from ic.infrastructure.retrieval.embedder import Embedder
from ic.infrastructure.retrieval.pgvector_store import PGVectorStore
from ic.infrastructure.sources.web_connector import WebConnector
from ic.shared.config import get_settings
from ic.shared.types import (
    CapBenchmarking,
    ChunkResult,
    ClasificacionGICS,
    DescripcionEmpresa,
    InformePiloto,
    Referencia,
)

settings = get_settings()


def _ollama_available() -> bool:
    try:
        s = socket.create_connection(("127.0.0.1", 11434), timeout=2)
        s.close()
        return True
    except OSError:
        return False


def _network_available() -> bool:
    try:
        s = socket.create_connection(("8.8.8.8", 53), timeout=2)
        s.close()
        return True
    except OSError:
        return False


requires_ollama = pytest.mark.skipif(not _ollama_available(), reason="Ollama no disponible")
requires_network = pytest.mark.skipif(not _network_available(), reason="Sin acceso a red")


# ---------------------------------------------------------------------------
# DocxBuilder — lógica pura
# ---------------------------------------------------------------------------


def test_docx_builder_creates_file(tmp_path):
    out = str(tmp_path / "test.docx")
    b = DocxBuilder()
    b.add_heading("Título de prueba", level=1)
    b.add_paragraph("Párrafo de prueba.")
    b.save(out)

    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert "Título de prueba" in texts
    assert "Párrafo de prueba." in texts


def test_docx_builder_headings_and_paragraphs(tmp_path):
    out = str(tmp_path / "headings.docx")
    b = DocxBuilder()
    b.add_heading("H1", level=1)
    b.add_heading("H2", level=2)
    b.add_paragraph("Normal text")
    b.save(out)

    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert "H1" in texts
    assert "H2" in texts
    assert "Normal text" in texts


def test_docx_builder_markdown_section(tmp_path):
    out = str(tmp_path / "markdown.docx")
    b = DocxBuilder()
    b.add_markdown_section("## Sección\n**negrita** y texto normal\n- item 1\n- item 2")
    b.save(out)

    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert any("Sección" in t for t in texts)
    assert any("item 1" in t for t in texts)


# ---------------------------------------------------------------------------
# APA Formatter — lógica pura
# ---------------------------------------------------------------------------


def test_apa_formatter_web_source_without_author():
    chunk = ChunkResult(
        chunk_texto="Contenido de ejemplo sobre el mercado alimentario.",
        fuente_tipo="web",
        fuente_url="https://www.fortune.com/alimentario/2024",
        score=0.9,
        metadatos={
            "titulo": "Market Report 2024",
            "fuente_nombre": "fortune.com",
            "año": "2024",
            "autor": None,
            "fecha_acceso": "2024-06-01",
        },
    )
    meta = extract_citation_metadata(chunk)
    inline = format_inline(meta)
    assert "fortune.com" in inline
    assert "2024" in inline

    ref = format_reference(meta)
    assert ref.año == "2024"
    assert "fortune.com" in ref.apa_string
    assert "Market Report 2024" in ref.apa_string


def test_apa_formatter_web_source_with_year():
    chunk = ChunkResult(
        chunk_texto="Datos de 2023 sobre competidores.",
        fuente_tipo="web",
        fuente_url="https://statista.com/report",
        score=0.85,
        metadatos={
            "titulo": "Industry Overview",
            "fuente_nombre": "statista.com",
            "año": "2023",
            "autor": None,
            "fecha_acceso": None,
        },
    )
    ref = format_reference(extract_citation_metadata(chunk))
    assert "s.f." not in ref.apa_string
    assert "2023" in ref.apa_string


def test_apa_formatter_deduplication():
    chunk_a = ChunkResult(
        chunk_texto="Texto A",
        fuente_tipo="web",
        fuente_url="https://example.com/page",
        score=0.9,
        metadatos={"titulo": "Page A", "fuente_nombre": "example.com", "año": "2022", "autor": None},
    )
    chunk_b = ChunkResult(
        chunk_texto="Texto B diferente",
        fuente_tipo="web",
        fuente_url="https://example.com/page",
        score=0.8,
        metadatos={"titulo": "Page A", "fuente_nombre": "example.com", "año": "2022", "autor": None},
    )
    refs = build_references_section([chunk_a, chunk_b])
    assert len(refs) == 1


# ---------------------------------------------------------------------------
# PilotAssembler — con datos fijos (sin Ollama)
# ---------------------------------------------------------------------------


def _make_informe() -> InformePiloto:
    empresa = DescripcionEmpresa(
        nombre_empresa="Empresa Ficticia S.A.",
        actividad_ciiu="1081",
        texto="**Empresa Ficticia** es una empresa de alimentos.\n- Galletas\n- Snacks",
    )
    gics = ClasificacionGICS(
        grupo_gics="Consumer Staples — Food Products",
        justificacion="Produce bienes de consumo básico.",
        texto_raw="Consumer Staples — Food Products",
    )
    chunks = [
        ChunkResult(
            chunk_texto="Competidor A lidera el mercado con 30% de participación.",
            fuente_tipo="web",
            fuente_url="https://fuente-a.com/informe",
            score=0.91,
            metadatos={"titulo": "Informe A", "fuente_nombre": "fuente-a.com", "año": "2024", "autor": None},
        ),
    ]
    cap = CapBenchmarking(
        texto="## Competidores\nLa empresa enfrenta competencia de varios actores.",
        n_chunks_usados=1,
        retrieval_suficiente=True,
        chunks_usados=chunks,
    )
    from ic.infrastructure.citation.apa_formatter import build_references_section
    refs = build_references_section(chunks)
    return InformePiloto(empresa=empresa, gics=gics, cap_benchmarking=cap, referencias=refs)


def test_assembler_with_mock_informe(tmp_path):
    out = str(tmp_path / "piloto.docx")
    builder = DocxBuilder()
    assembler = PilotAssembler(builder)

    informe = _make_informe()
    result_path = assembler.assemble(informe, out)

    assert result_path == out
    doc = Document(out)
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "Empresa Ficticia" in all_text
    assert "Consumer Staples" in all_text
    assert "Competidores" in all_text
    assert "fuente-a.com" in all_text


# ---------------------------------------------------------------------------
# End-to-end con Ollama + BD real
# ---------------------------------------------------------------------------


@pytest.mark.slow
@requires_ollama
@requires_network
def test_pilot_end_to_end_noel(tmp_path):
    from ic.data.database import SessionLocal
    from ic.data.repositories.embedding_repo import EmbeddingRepository

    out = str(tmp_path / "noel_piloto.docx")

    llm = build_llm(settings)
    embedder = Embedder(settings)

    with SessionLocal() as db:
        repo = EmbeddingRepository(db)
        store = PGVectorStore(embedder, repo, settings)
        web = WebConnector(settings)

        uc_desc = GenerateEmpresaDescription(chain=build_chain_01(llm), web_connector=web)
        empresa = uc_desc.execute("Noel S.A.S.")

        uc_gics = ClassifyGICS(chain=build_chain_02(llm))
        gics = uc_gics.execute(nombre_empresa="Noel S.A.S.", descripcion_empresa=empresa.texto)

        uc_bench = GenerateBenchmarking(
            chain=build_chain_03(llm),
            store=store,
            web_connector=web,
        )
        cap = uc_bench.execute(
            nombre_empresa="Noel S.A.S.",
            descripcion_empresa=empresa.texto,
            sector_gics=gics.grupo_gics,
        )

    from ic.infrastructure.citation.apa_formatter import build_references_section

    refs = build_references_section(cap.chunks_usados)

    informe = InformePiloto(
        empresa=empresa,
        gics=gics,
        cap_benchmarking=cap,
        referencias=refs,
    )

    assembler = PilotAssembler(DocxBuilder())
    result_path = assembler.assemble(informe, out)

    doc = Document(result_path)
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert len(doc.paragraphs) > 5
    assert "Noel" in all_text or "noel" in all_text.lower()
