"""T16.2 + T16.3: tests de estilo Normal y tablas Word en DocxRenderer."""

from __future__ import annotations

from docx import Document as WDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt

from ic.domain.report_model import ReportDocument, TableBlock
from ic.infrastructure.render.docx_renderer import NORMAL_FONT_NAME, DocxRenderer


def _empty_doc() -> ReportDocument:
    return ReportDocument(title="Test", company="Empresa", date="2026-01-01", blocks=[])


def test_normal_style_font_size(tmp_path):
    """Estilo Normal tiene fuente Open Sans 11pt (T17.4, DT-024)."""
    path = str(tmp_path / "out.docx")
    DocxRenderer().render(_empty_doc(), path)
    word = WDocument(path)
    style = word.styles["Normal"]
    assert style.font.name == NORMAL_FONT_NAME == "Open Sans"
    assert style.font.size == Pt(11)


def test_normal_style_line_spacing(tmp_path):
    """Estilo Normal tiene interlineado múltiple 1,16."""
    path = str(tmp_path / "out.docx")
    DocxRenderer().render(_empty_doc(), path)
    word = WDocument(path)
    pf = word.styles["Normal"].paragraph_format
    assert pf.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert abs(pf.line_spacing - 1.16) < 0.01


def test_normal_style_justify(tmp_path):
    """Estilo Normal tiene alineación justificada."""
    path = str(tmp_path / "out.docx")
    DocxRenderer().render(_empty_doc(), path)
    word = WDocument(path)
    pf = word.styles["Normal"].paragraph_format
    assert pf.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_normal_style_space_after(tmp_path):
    """Estilo Normal tiene espacio después de 8pt."""
    path = str(tmp_path / "out.docx")
    DocxRenderer().render(_empty_doc(), path)
    word = WDocument(path)
    pf = word.styles["Normal"].paragraph_format
    assert pf.space_after == Pt(8)


def test_normal_style_widow_control(tmp_path):
    """Estilo Normal tiene control de viudas/huérfanas activado."""
    path = str(tmp_path / "out.docx")
    DocxRenderer().render(_empty_doc(), path)
    word = WDocument(path)
    pf = word.styles["Normal"].paragraph_format
    assert pf.widow_control is True


# ---------------------------------------------------------------------------
# T16.3: tabla Word de ficha de competidor
# ---------------------------------------------------------------------------


def test_docx_table_block_rendered(tmp_path):
    """Un TableBlock en el ReportDocument se renderiza como tabla Word nativa."""
    block = TableBlock(
        headers=("Campo", "Valor"),
        rows=(
            ("Nombre + País", "Nestlé — Suiza"),
            ("Estrategia", "Líder global en nutrición"),
        ),
    )
    doc = ReportDocument(title="T", company="C", date="2026", blocks=[block])
    path = str(tmp_path / "tabla.docx")
    DocxRenderer().render(doc, path)
    word = WDocument(path)

    assert len(word.tables) >= 1, "El .docx debe contener al menos 1 tabla Word"
    tbl = word.tables[0]
    assert tbl.rows[0].cells[0].text == "Campo"
    assert tbl.rows[0].cells[1].text == "Valor"
    assert tbl.rows[1].cells[0].text == "Nombre + País"
    assert tbl.rows[2].cells[1].text == "Líder global en nutrición"


def test_docx_multiple_tables_rendered(tmp_path):
    """Múltiples TableBlocks generan múltiples tablas Word."""
    blocks = [
        TableBlock(headers=("A", "B"), rows=(("a1", "b1"),)),
        TableBlock(headers=("X", "Y"), rows=(("x1", "y1"),)),
    ]
    doc = ReportDocument(title="T", company="C", date="2026", blocks=blocks)
    path = str(tmp_path / "multi.docx")
    DocxRenderer().render(doc, path)
    word = WDocument(path)
    assert len(word.tables) >= 2
