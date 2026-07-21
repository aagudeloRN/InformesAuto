"""Tests del DocxRenderer: portada, TOC, tablas, cursivas APA, figuras, callouts."""

from pathlib import Path

from docx import Document

from ic.domain.report_model import (
    Callout,
    CitationBlock,
    FigureBlock,
    Heading,
    PageBreak,
    Paragraph,
    ReportDocument,
    TableBlock,
    Toc,
)
from ic.infrastructure.render.docx_renderer import DocxRenderer


def _make_doc(**kwargs) -> ReportDocument:
    defaults = {
        "title": "Informe IC — Empresa Test",
        "company": "Empresa Test S.A.",
        "date": "21 de julio de 2026",
        "blocks": [],
    }
    defaults.update(kwargs)
    return ReportDocument(**defaults)


def _all_text(doc: Document) -> str:
    return " ".join(p.text for p in doc.paragraphs)


def _minimal_png(path: Path) -> None:
    """Crea un PNG válido de 10×10 píxel usando Pillow."""
    from PIL import Image
    img = Image.new("RGB", (10, 10), color=(0, 128, 255))
    img.save(str(path), format="PNG")


# ---------------------------------------------------------------------------
# Portada
# ---------------------------------------------------------------------------


def test_cover_page_contains_company(tmp_path):
    out = str(tmp_path / "out.docx")
    doc = _make_doc()
    DocxRenderer().render(doc, out)
    word_doc = Document(out)
    text = _all_text(word_doc)
    assert "Empresa Test S.A." in text


def test_cover_page_contains_date(tmp_path):
    out = str(tmp_path / "out.docx")
    DocxRenderer().render(_make_doc(), out)
    text = _all_text(Document(out))
    assert "21 de julio de 2026" in text


# ---------------------------------------------------------------------------
# TOC
# ---------------------------------------------------------------------------


def test_toc_field_present(tmp_path):
    out = str(tmp_path / "out.docx")
    doc = _make_doc(blocks=[Toc()])
    DocxRenderer().render(doc, out)
    word_doc = Document(out)
    # El campo TOC queda como instrucción XML; verificamos que "Tabla de Contenidos" aparece
    text = _all_text(word_doc)
    assert "Tabla de Contenidos" in text


# ---------------------------------------------------------------------------
# Headings y paragraphs
# ---------------------------------------------------------------------------


def test_heading_rendered(tmp_path):
    out = str(tmp_path / "out.docx")
    blocks = [Heading("Capítulo 2. Análisis Competitivo", level=2)]
    DocxRenderer().render(_make_doc(blocks=blocks), out)
    text = _all_text(Document(out))
    assert "Capítulo 2" in text


def test_inline_bold_and_italic(tmp_path):
    out = str(tmp_path / "out.docx")
    blocks = [Paragraph("Texto con **negrita** y *cursiva* aquí.")]
    DocxRenderer().render(_make_doc(blocks=blocks), out)
    word_doc = Document(out)
    # Verificar que los runs contienen el texto sin marcadores
    all_runs = [r for p in word_doc.paragraphs for r in p.runs]
    run_texts = [r.text for r in all_runs]
    assert any("negrita" in t for t in run_texts)
    assert any("cursiva" in t for t in run_texts)
    bold_runs = [r for r in all_runs if r.bold and r.text]
    italic_runs = [r for r in all_runs if r.italic and r.text]
    assert any("negrita" in r.text for r in bold_runs)
    assert any("cursiva" in r.text for r in italic_runs)


# ---------------------------------------------------------------------------
# Tabla
# ---------------------------------------------------------------------------


def test_table_block_renders_correctly(tmp_path):
    out = str(tmp_path / "out.docx")
    block = TableBlock(
        headers=("Empresa", "País", "Cuota"),
        rows=(("Acme Corp", "USA", "30%"), ("Beta GmbH", "DE", "15%")),
        caption="Tabla 1. Competidores",
    )
    DocxRenderer().render(_make_doc(blocks=[block]), out)
    word_doc = Document(out)
    assert len(word_doc.tables) == 1
    table = word_doc.tables[0]
    assert table.rows[0].cells[0].text == "Empresa"
    assert table.rows[1].cells[0].text == "Acme Corp"
    assert table.rows[2].cells[2].text == "15%"
    text = _all_text(word_doc)
    assert "Tabla 1. Competidores" in text


# ---------------------------------------------------------------------------
# CitationBlock — sangría francesa + cursivas APA
# ---------------------------------------------------------------------------


def test_citation_italic_apa(tmp_path):
    out = str(tmp_path / "out.docx")
    apa = "Autor, A. (2024). *Título del artículo*. *Revista Académica*, 10(2), 45–60."
    DocxRenderer().render(_make_doc(blocks=[CitationBlock(apa)]), out)
    word_doc = Document(out)
    italic_runs = [
        r for p in word_doc.paragraphs for r in p.runs if r.italic and r.text.strip()
    ]
    assert len(italic_runs) >= 1, "Debe haber al menos un run en cursiva en la referencia APA"
    italic_text = " ".join(r.text for r in italic_runs)
    assert "Título del artículo" in italic_text or "Revista Académica" in italic_text


# ---------------------------------------------------------------------------
# FigureBlock
# ---------------------------------------------------------------------------


def test_figure_block_missing_image(tmp_path):
    out = str(tmp_path / "out.docx")
    block = FigureBlock(
        image_path="/no/existe.png",
        alt="Radar benchmarking",
        caption="Figura 1",
    )
    DocxRenderer().render(_make_doc(blocks=[block]), out)
    text = _all_text(Document(out))
    assert "Radar benchmarking" in text or "no disponible" in text


def test_figure_block_with_image(tmp_path):
    png_path = tmp_path / "chart.png"
    _minimal_png(png_path)
    out = str(tmp_path / "out.docx")
    block = FigureBlock(
        image_path=str(png_path),
        alt="Radar benchmarking competidores",
        caption="Figura 1. Posicionamiento",
    )
    DocxRenderer().render(_make_doc(blocks=[block]), out)
    word_doc = Document(out)
    # La imagen embebida aparece como InlineShape
    shapes = word_doc.inline_shapes
    assert len(shapes) == 1
    text = _all_text(word_doc)
    assert "Figura 1. Posicionamiento" in text


# ---------------------------------------------------------------------------
# Callout
# ---------------------------------------------------------------------------


def test_callout_absent_data_in_text(tmp_path):
    out = str(tmp_path / "out.docx")
    block = Callout(
        text="Información no disponible: fuentes insuficientes para este capítulo.",
        kind="absent_data",
    )
    DocxRenderer().render(_make_doc(blocks=[block]), out)
    text = _all_text(Document(out))
    assert "Información no disponible" in text


def test_callout_actionable_in_text(tmp_path):
    out = str(tmp_path / "out.docx")
    block = Callout(text="Accionables requieren validación directiva.", kind="actionable")
    DocxRenderer().render(_make_doc(blocks=[block]), out)
    text = _all_text(Document(out))
    assert "validación directiva" in text


# ---------------------------------------------------------------------------
# PageBreak
# ---------------------------------------------------------------------------


def test_page_break_does_not_crash(tmp_path):
    out = str(tmp_path / "out.docx")
    blocks = [Heading("Cap 1", level=2), PageBreak(), Heading("Cap 2", level=2)]
    DocxRenderer().render(_make_doc(blocks=blocks), out)
    text = _all_text(Document(out))
    assert "Cap 1" in text and "Cap 2" in text
