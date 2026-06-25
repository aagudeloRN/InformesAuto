import os
import re

from docx import Document
from docx.shared import Pt


class DocxBuilder:
    def __init__(self, template_path: str | None = None) -> None:
        if template_path and os.path.exists(template_path):
            self._doc = Document(template_path)
        else:
            self._doc = Document()

    def add_heading(self, text: str, level: int = 1) -> None:
        self._doc.add_heading(text, level=level)

    def add_paragraph(self, text: str, style: str = "Normal") -> None:
        self._doc.add_paragraph(text, style=style)

    def add_markdown_section(self, markdown_text: str) -> None:
        for line in markdown_text.splitlines():
            heading_match = re.match(r"^(#{1,6})\s+(.*)", line)
            if heading_match:
                level = min(len(heading_match.group(1)), 4)
                self._doc.add_heading(heading_match.group(2).strip(), level=level)
                continue

            if line.startswith("- ") or line.startswith("* "):
                text = line[2:].strip()
                text = _strip_inline_bold(text)
                self._doc.add_paragraph(text, style="List Bullet")
                continue

            stripped = line.strip()
            if not stripped:
                continue

            para = self._doc.add_paragraph()
            _add_inline_bold(para, stripped)

    def save(self, output_path: str) -> None:
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        self._doc.save(output_path)


def _strip_inline_bold(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def _add_inline_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        bold_match = re.match(r"\*\*(.+?)\*\*", part)
        if bold_match:
            run = paragraph.add_run(bold_match.group(1))
            run.bold = True
        else:
            paragraph.add_run(part)
